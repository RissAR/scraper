import logging
from pathlib import Path
from typing import Sequence

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from models import ProductItem

logger = logging.getLogger(__name__)


def export_items_to_excel(
    items: Sequence[ProductItem], filename: str = "export.xlsx"
) -> str:
    """
    Експортує список ProductItem у стилізований Excel-файл (.xlsx).
    Зберігає файл у папку 'data/' та повертає абсолютний або відносний шлях до нього.
    """
    export_dir = Path("data")
    export_dir.mkdir(parents=True, exist_ok=True)
    file_path = export_dir / filename

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Товари"

    headers = [
        "Платформа",
        "ID",
        "Назва",
        "Ціна (грн)",
        "Місто/Магазин",
        "Реклама (ТОП)",
        "Посилання",
        "Дата збору",
    ]
    ws.append(headers)

    # Стилізація шапки
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(
        start_color="1F4E79", end_color="1F4E79", fill_type="solid"
    )
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # Заповнення даними
    data_font = Font(name="Calibri", size=10)
    link_font = Font(name="Calibri", size=10, color="0563C1", underline="single")
    default_alignment = Alignment(vertical="center")

    for row_idx, item in enumerate(items, start=2):
        platform_str = (
            item.platform.value.upper()
            if hasattr(item.platform, "value")
            else str(item.platform).upper()
        )
        is_promoted_str = "Так" if item.is_promoted else "Ні"
        scraped_date_str = item.scraped_at.strftime("%Y-%m-%d %H:%M:%S")

        ws.cell(row=row_idx, column=1, value=platform_str).alignment = Alignment(
            horizontal="center", vertical="center"
        )
        ws.cell(row=row_idx, column=2, value=item.id).alignment = Alignment(
            horizontal="center", vertical="center"
        )
        ws.cell(row=row_idx, column=3, value=item.title).alignment = default_alignment
        ws.cell(row=row_idx, column=4, value=item.price).alignment = Alignment(
            horizontal="right", vertical="center"
        )
        ws.cell(row=row_idx, column=5, value=item.location or "-").alignment = (
            default_alignment
        )
        ws.cell(row=row_idx, column=6, value=is_promoted_str).alignment = Alignment(
            horizontal="center", vertical="center"
        )

        url_cell = ws.cell(row=row_idx, column=7, value=item.url)
        url_cell.hyperlink = item.url
        url_cell.font = link_font
        url_cell.alignment = default_alignment

        ws.cell(row=row_idx, column=8, value=scraped_date_str).alignment = Alignment(
            horizontal="center", vertical="center"
        )

        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=row_idx, column=col_idx)
            if col_idx != 7:
                c.font = data_font

    # Автопідбір ширини колонок
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val = str(cell.value or "")
            if len(val) > max_len:
                max_len = len(val)
        # Обмежуємо ширину між 12 та 50 символами
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    wb.save(file_path)
    logger.info("Successfully exported %d items to Excel: %s", len(items), file_path)
    return str(file_path)
